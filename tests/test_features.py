import io
import os
import re
import tempfile
import unittest

from openpyxl import load_workbook

from portal import create_app
from portal.core import dataset as ds
from portal.core.analytics import compare_models, holt_damped
from portal.core.custom import analyse, hathras_example
from portal.core.geo import haversine_km
from portal.scenario import hathras_scenario


class EngineFeatureTests(unittest.TestCase):
    def test_holt_on_straight_line_extends_it(self):
        out = holt_damped([10, 20, 30, 40, 50], 2)
        self.assertGreater(out["rows"][0]["value"], 50)
        self.assertGreater(out["rows"][1]["value"], out["rows"][0]["value"])
        self.assertLess(out["rows"][0]["lower"], out["rows"][0]["value"])

    def test_model_comparison_uses_rolling_evaluation(self):
        cm = compare_models(ds.PRODUCTION_SERIES)
        self.assertEqual(cm["primary"], "linear")
        ev = cm["evaluation"]
        self.assertEqual(ev["assessment"]["level"], "indicative")  # 3 one-step tests only
        self.assertEqual(ev["assessment"]["points"], 3)

    def test_whatif_levers(self):
        base = hathras_scenario({})
        shock = hathras_scenario({"yield_shock": -0.12, "extra_capacity": 50000, "tariff_change": 0.15})
        self.assertAlmostEqual(shock["production_t"], base["production_t"] * 0.88)
        self.assertEqual(shock["inputs"]["total_capacity"], 1_000_000)
        self.assertLess(shock["simulation"]["total_overflow_t"], base["simulation"]["total_overflow_t"])
        self.assertAlmostEqual(shock["economics"]["rate_rs_kwh"], 7.5 * 1.15)
        self.assertGreater(base["economics"]["season_cost_rs"], 0)

    def test_block_monthly_and_nearest(self):
        sc = hathras_scenario({})
        b = next(x for x in sc["blocks"] if x["id"] == "sikandra-rao")
        peak = max(m["utilisation"] for m in b["monthly"])
        self.assertAlmostEqual(peak, b["utilisation"])
        self.assertEqual(b["monthly"][0]["utilisation"], 0)  # nothing stored in January
        for n in b["nearest_headroom"]:
            self.assertGreater(n["headroom_t"], 0)
        self.assertAlmostEqual(haversine_km({"lat": 0, "lng": 0}, {"lat": 0, "lng": 1}), 111.19, places=1)

    def test_custom_analysis_carries_whatifs(self):
        p = hathras_example()
        p.update(yield_shock=-0.1, extra_capacity=100000)
        a = analyse(p)
        self.assertEqual(a["inputs"]["total_capacity"], 1_050_000)
        self.assertIn("smoothing", a["forecast_compare"])


class FeatureRouteTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        os.environ.pop("VERCEL", None)
        self.app = create_app({"TESTING": True, "SQLITE_PATH": os.path.join(self.tmp.name, "t.db"), "DATABASE_URL": ""})
        self.c = self.app.test_client()

    def tearDown(self):
        self.tmp.cleanup()

    def test_templates_import_export(self):
        x = self.c.get("/templates/coldflow-template.xlsx")
        self.assertEqual(x.status_code, 200)
        r = self.c.post("/api/import", data={"file": (io.BytesIO(x.data), "t.xlsx")}, content_type="multipart/form-data")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(len(r.get_json()["series"]), 3)
        self.assertEqual(r.get_json()["settings"]["capacity"], 100000)
        csv_body = b"season,area_ha,production_t\n2021-22,3200,81600\n"
        r = self.c.post("/api/import", data={"file": (io.BytesIO(csv_body), "a.csv")}, content_type="multipart/form-data")
        self.assertEqual(r.get_json()["series"][0]["production_t"], 81600)
        bad = self.c.post("/api/import", data={"file": (io.BytesIO(b"x"), "a.pdf")}, content_type="multipart/form-data")
        self.assertEqual(bad.status_code, 422)
        e = self.c.post("/api/export", json={"kind": "hathras", "params": {"extra_capacity": 50000}})
        self.assertEqual(e.status_code, 200)
        wb = load_workbook(io.BytesIO(e.data))
        self.assertEqual(wb.sheetnames, ["Summary", "Seasons", "Forecast", "Evaluation", "Monthly storage", "Blocks", "Assumptions"])
        bad = hathras_example(); bad["series"] = []
        self.assertEqual(self.c.post("/api/export", json={"kind": "plan", "payload": bad}).status_code, 422)

    def test_workspaces_districts_pwa(self):
        for role in ("owner", "fpo", "admin"):
            with self.subTest(role=role):
                self.assertEqual(self.c.get(f"/workspace/{role}").status_code, 200)
        self.assertEqual(self.c.get("/workspace/nope").status_code, 404)
        self.assertEqual(self.c.get("/api/districts/hathras").get_json()["series"][0]["season"], "2018-19")
        self.assertEqual(self.c.get("/api/districts/agra").status_code, 404)
        sw = self.c.get("/sw.js")
        self.assertEqual(sw.headers["Service-Worker-Allowed"], "/")
        v = self.app.config["STATIC_VERSION"]
        body = sw.get_data(as_text=True)
        self.assertIn(f"coldflow-{v}", body)
        self.assertNotIn("{{", body)
        # Pages reference content-versioned assets, so a deploy busts browser and SW caches.
        self.assertIn(f"/static/css/site.css?v={v}", self.c.get("/").get_data(as_text=True))
        self.assertIn("standalone", self.c.get("/manifest.webmanifest").get_data(as_text=True))
        self.assertEqual(self.c.get("/offline").status_code, 200)

    def test_role_on_signup_shows_workspace(self):
        page = self.c.get("/signup").get_data(as_text=True)
        tok = re.search(r'name="csrf_token" value="([^"]+)"', page).group(1)
        self.c.post("/signup", data={"csrf_token": tok, "name": "Ops Person", "email": "ops@example.test", "role": "admin",
                                     "password": "test-pass-123", "confirm": "test-pass-123", "terms": "1"})
        dash = self.c.get("/dashboard").get_data(as_text=True)
        self.assertIn("District administrator workspace", dash)
        self.assertIn("Produce at spoilage risk", dash)


if __name__ == "__main__":
    unittest.main()

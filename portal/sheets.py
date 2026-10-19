"""CSV/Excel templates, file import, and Excel report export."""
from __future__ import annotations

import csv
import io
from datetime import datetime, timezone

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .core import dataset as ds
from .core.analytics import compare_models

TEMPLATE_HEADER = ["season", "area_ha", "production_t"]
EXAMPLE_ROWS = [["2021-22", 3200, 81600], ["2022-23", 3350, 86100], ["2023-24", 3420, 88900]]
SETTING_KEYS = {  # label in the Settings sheet -> payload key
    "cold storage capacity (t)": "capacity",
    "share of harvest stored (0-1)": "ratio",
    "alert threshold (0-1)": "threshold",
    "one facility's capacity (t)": "facility",
}
MAX_ROWS = 60

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1C5CAB")
TITLE_FONT = Font(bold=True, size=14)


class ImportError_(ValueError):
    pass


# ---------------------------------------------------------------- templates

def template_csv() -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(TEMPLATE_HEADER)
    w.writerows(EXAMPLE_ROWS)
    return buf.getvalue()


def _style_header(ws, row=1):
    for cell in ws[row]:
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL


def _autosize(ws, min_w=10, max_w=48):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(min_w, min(max_w, width + 2))


def template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Seasons"
    ws.append(TEMPLATE_HEADER)
    _style_header(ws)
    for r in EXAMPLE_ROWS:
        ws.append(r)
    _autosize(ws, 14)

    st = wb.create_sheet("Settings")
    st.append(["setting", "value"])
    _style_header(st)
    for label, key in SETTING_KEYS.items():
        default = {"capacity": 100000, "ratio": ds.DEFAULT_STORAGE_RATIO, "threshold": ds.DEFAULT_ALERT_THRESHOLD,
                   "facility": 10000}[key]
        st.append([label, default])
    _autosize(st, 14)

    info = wb.create_sheet("Instructions")
    for line in [
        "ColdFlow import template",
        "",
        "Seasons sheet: one row per season, oldest first. Season like 2023-24, area in hectares, production in tonnes.",
        "Use at least 3 consecutive seasons. Replace the example rows with your own figures.",
        "Settings sheet (optional): your storage capacity, share of harvest stored, alert threshold and one facility's size.",
        "Upload this file on ColdFlow's Plan page (Import CSV / Excel).",
    ]:
        info.append([line])
    info["A1"].font = TITLE_FONT
    info.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- import

def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _rows_to_series(rows):
    """First row may be a header; finds season/area/production columns by name or position."""
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        raise ImportError_("The file has no rows.")
    head = [str(c or "").strip().lower() for c in rows[0]]
    idx = {"season": 0, "area": 1, "prod": 2}
    if any(k in h for h in head for k in ("season", "year", "area", "prod")):
        for i, h in enumerate(head):
            if "season" in h or "year" in h:
                idx["season"] = i
            elif "area" in h:
                idx["area"] = i
            elif "prod" in h or "output" in h or "harvest" in h:
                idx["prod"] = i
        rows = rows[1:]
    series = []
    for r in rows[:MAX_ROWS]:
        r = list(r) + [None] * 3
        season = r[idx["season"]]
        if isinstance(season, datetime):
            season = f"{season.year}-{(season.year + 1) % 100:02d}"
        series.append({"season": str(season or "").strip(), "area_ha": _num(r[idx["area"]]),
                       "production_t": _num(r[idx["prod"]])})
    return series


def parse_upload(filename: str, data: bytes) -> dict:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise ImportError_("That Excel file couldn't be read. Save it as .xlsx and try again.")
        ws = wb["Seasons"] if "Seasons" in wb.sheetnames else wb.worksheets[0]
        series = _rows_to_series(list(ws.iter_rows(values_only=True)))
        settings = {}
        if "Settings" in wb.sheetnames:
            for row in wb["Settings"].iter_rows(values_only=True):
                if row and row[0] and str(row[0]).strip().lower() in SETTING_KEYS and len(row) > 1:
                    v = _num(row[1])
                    if v is not None:
                        settings[SETTING_KEYS[str(row[0]).strip().lower()]] = v
        return {"series": series, "settings": settings}
    if name.endswith((".csv", ".txt")) or not name:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        try:
            dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return {"series": _rows_to_series(list(csv.reader(io.StringIO(text), dialect))), "settings": {}}
    raise ImportError_("Upload a .csv or .xlsx file.")


# ---------------------------------------------------------------- export

def _sheet_table(wb, title, header, rows, number_formats=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    _style_header(ws)
    for r in rows:
        ws.append(r)
    for col, fmt in (number_formats or {}).items():
        for cell in ws[col][1:]:
            cell.number_format = fmt
    _autosize(ws)
    ws.freeze_panes = "A2"
    return ws


def report_xlsx(title: str, analysis: dict, *, blocks=None) -> bytes:
    """One workbook: Summary, Seasons, Forecast (both models), Monthly storage, Blocks, Assumptions."""
    sim, i, econ = analysis["simulation"], analysis["inputs"], analysis.get("economics") or {}
    series = analysis.get("series") or []
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"ColdFlow report · {title}"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated {datetime.now(timezone.utc):%d %b %Y %H:%M} UTC · estimates for planning, not official statistics"
    status = {"good": "Adequate", "warning": "Watch", "serious": "Strained", "critical": "Deficit"}[sim["status"]]
    rows = [
        ("Season in outlook", i["season"]),
        ("Status", status),
        ("Harvest (t)", analysis["production_t"]),
        ("Yield shock applied", i.get("yield_shock", 0)),
        ("Storage demand (t)", analysis["demand_t"]),
        ("Cold storage capacity (t)", i.get("total_capacity", i["capacity"])),
        ("  of which new capacity (t)", i.get("extra_capacity", 0)),
        ("Peak utilisation", sim["peak_utilisation"]),
        ("Peak month", sim["peak_month"]),
        ("Deficit months", ", ".join(sim["deficit_months"]) or "None"),
        ("Without cold storage (t)", sim["total_overflow_t"]),
        ("Capacity needed to avoid deficit (t)", sim["capacity_needed_t"]),
        ("Capacity needed under alert threshold (t)", sim["capacity_needed_threshold_t"]),
        ("Holding power cost for the season (₹)", econ.get("season_cost_rs")),
        ("Power cost per tonne stored (₹)", econ.get("cost_per_t_stored_rs")),
        ("Power tariff used (₹/kWh)", econ.get("rate_rs_kwh")),
    ]
    for r, (k, v) in enumerate(rows, start=4):
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        c = ws.cell(row=r, column=2, value=v)
        if k in ("Peak utilisation", "Yield shock applied"):
            c.number_format = "0.0%"
        elif isinstance(v, float):
            c.number_format = "#,##0"
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 22
    for row in ws.iter_rows(min_row=4, max_col=2):
        row[1].alignment = Alignment(horizontal="right")

    if series:
        _sheet_table(wb, "Seasons", ["Season", "Area (ha)", "Production (t)", "Yield (kg/ha)", "Output YoY"],
                     [[r["season"], r["area_ha"], r["production_t"], r["yield_kg_ha"], r.get("production_yoy")] for r in series],
                     {"B": "#,##0", "C": "#,##0", "D": "#,##0", "E": "0.00%"})
        cm = compare_models(series)
        frows = []
        for name, key in (("Linear trend", "linear"), ("Exponential smoothing", "smoothing")):
            for r in cm[key]["rows"]:
                frows.append([name, r["season"], r["production_t"], r["production_lower"], r["production_upper"], r["yield_kg_ha"]])
        _sheet_table(wb, "Forecast", ["Model", "Season", "Production (t)", "95% low", "95% high", "Yield (kg/ha)"], frows,
                     {"C": "#,##0", "D": "#,##0", "E": "#,##0", "F": "#,##0"})

    _sheet_table(wb, "Monthly storage", ["Month", "Intake (t)", "Release (t)", "Stock held (t)", "Unmet (t)", "Demand vs capacity"],
                 [[m["month"], m["intake"], m["release"], m["stock"], m["unmet"], m["utilisation"]] for m in sim["months"]],
                 {"B": "#,##0", "C": "#,##0", "D": "#,##0", "E": "#,##0", "F": "0.0%"})
    if blocks:
        _sheet_table(wb, "Blocks", ["Block", "Demand (t)", "Capacity (t)", "Gap (t)", "Utilisation"],
                     [[b["name"], b["demand_t"], b["capacity_t"], b["gap_t"], b["utilisation"]] for b in blocks],
                     {"B": "#,##0", "C": "#,##0", "D": "#,##0", "E": "0%"})
    _sheet_table(wb, "Assumptions", ["Item", "Value", "Status", "Note"],
                 [[a["item"], a["value"], a["status"], a["note"]] for a in ds.ASSUMPTIONS])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
